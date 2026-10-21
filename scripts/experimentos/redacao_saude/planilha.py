"""Monta a planilha do experimento "redacao de saude" a partir dos resultados baixados do S3.

    python scripts/experimentos/redacao_saude/rodar_na_aws.py baixar --data AAAA-MM-DD --para <pasta>
    python scripts/experimentos/redacao_saude/planilha.py --resultados <pasta> --perguntas \
        scripts/experimentos/redacao_saude/perguntas.yaml --xlsx <arquivo.xlsx>

Abas: Resumo, Por categoria, Respostas, Amostra para o dono (30 perguntas, respostas A-D embaralhadas por
linha, chave em colunas ocultas), Perguntas do dono (a resposta ATUAL reconstruida e as 4 novas).
Precisa de `openpyxl` (nao e' dependencia do projeto: instale no Python que roda este script).
"""

from __future__ import annotations

import argparse
import json
import random
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from experimento import (  # noqa: E402
    JUIZ,
    JUIZ2,
    MODELOS,
    _ler_jsonl,
    amostra,
    analisar,
    ler_perguntas,
)

NOMES_CRITERIOS = {
    "media_utilidade": "Utilidade (1-5)",
    "media_exatidao": "Exatidão médica (1-5)",
    "media_naturalidade": "Naturalidade/tom (1-5)",
    "pct_acertou_assunto": "% acertou o assunto",
    "pct_seguiu_esperado": "% seguiu o comportamento esperado",
    "pct_conduta_indevida": "% conduta indevida (juiz)",
    "pct_usou_o_que_disse": "% usou o que a pessoa disse (quando aplicável)",
    "pct_tamanho_adequado": "% tamanho adequado",
    "pct_fonte_incorreta": "% fonte citada incorreta",
    "pct_primeira": "% passou de primeira na verificação",
    "pct_reescrita": "% precisou reescrever",
    "pct_padrao": "% caiu na resposta padrão",
    "latencia_p50_s": "Latência p50 (s, até a resposta final)",
    "latencia_p95_s": "Latência p95 (s)",
    "latencia_1a_chamada_p50_s": "Latência p50 da 1a chamada (s)",
    "custo_usd_por_mil": "Custo estimado US$ por 1.000 respostas (só geração)",
    "palavras_media": "Palavras por resposta (média)",
    "erros_de_geracao": "Erros de geração (após 3 repetições)",
    "julgadas": "Respostas julgadas",
}


def _carregar(pasta: Path) -> dict[str, Any]:
    def ler(nome: str) -> str | None:
        caminho = pasta / nome
        return caminho.read_text(encoding="utf-8") if caminho.exists() else None

    return {
        "manifesto": json.loads(ler("manifesto.json") or "{}"),
        "respostas": _ler_jsonl(ler("respostas.jsonl")),
        "juiz1": {v["chave"]: v for v in _ler_jsonl(ler("juiz_opus.jsonl"))},
        "juiz2": {v["chave"]: v for v in _ler_jsonl(ler("juiz_glm.jsonl"))},
        "atual": json.loads(ler("atual.json") or "{}"),
    }


def montar(pasta: Path, perguntas_yaml: Path, xlsx: Path) -> dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    perguntas = ler_perguntas(perguntas_yaml.read_text(encoding="utf-8"))
    d = _carregar(pasta)
    modelos = [m for m in MODELOS if any(r["modelo"] == m for r in d["respostas"])]
    resumo = analisar(perguntas, d["respostas"], d["juiz1"], d["juiz2"], modelos)
    resp = {f"{r['pergunta_id']}|{r['modelo']}": r for r in d["respostas"]}
    j1 = d["juiz1"]

    negrito = Font(bold=True)
    cabec = PatternFill("solid", fgColor="DDE6F0")
    quebra = Alignment(wrap_text=True, vertical="top")

    def cabecalho(ws: Any, linha: int, valores: list[Any]) -> None:
        for i, v in enumerate(valores, start=1):
            c = ws.cell(row=linha, column=i, value=v)
            c.font = negrito
            c.fill = cabec
            c.alignment = quebra

    def larguras(ws: Any, mapa: dict[int, float]) -> None:
        for col, w in mapa.items():
            ws.column_dimensions[get_column_letter(col)].width = w

    wb = Workbook()

    # --- Resumo ---
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "Experimento: a Helena REDIGE respostas de saúde (vs. textos fixos da base) — dev, 17/10/2026"
    ws["A1"].font = Font(bold=True, size=13)
    man = d["manifesto"]
    ws["A2"] = (
        f"100 perguntas sintéticas; prompt {man.get('prompt_versao', '?')}; commit {man.get('git_sha', '?')[:8]}; "
        f"juiz {JUIZ['id']} (todas) e {JUIZ2['id']} (amostra de {resumo['concordancia']['pares']} pares)."
    )
    linha = 4
    cabecalho(ws, linha, ["Critério", *modelos])
    for chave, nome in NOMES_CRITERIOS.items():
        linha += 1
        ws.cell(row=linha, column=1, value=nome)
        for i, m in enumerate(modelos, start=2):
            ws.cell(row=linha, column=i, value=resumo["por_modelo"][m].get(chave))
    linha += 2
    ws.cell(
        row=linha, column=1, value="Concordância entre juízes (Opus 5 x GLM-5, mesma amostra)"
    ).font = negrito
    linha += 1
    cabecalho(
        ws,
        linha,
        [
            "Critério",
            "% igual / concorda",
            "% diferença ≤ 1",
            "Média Opus / % sim Opus",
            "Média GLM / % sim GLM",
            "Pearson / kappa",
        ],
    )
    conc = resumo["concordancia"]
    for c in ("utilidade", "exatidao", "naturalidade"):
        linha += 1
        x = conc[c]
        for i, v in enumerate(
            [c, x["pct_igual"], x["pct_dif_ate_1"], x["media_opus"], x["media_glm"], x["pearson"]], start=1
        ):
            ws.cell(row=linha, column=i, value=v)
    for c in ("acertou_assunto", "seguiu_esperado", "conduta_indevida", "tamanho_adequado"):
        linha += 1
        x = conc[c]
        for i, v in enumerate(
            [c, x["pct_concorda"], None, x["pct_true_opus"], x["pct_true_glm"], x["kappa"]], start=1
        ):
            ws.cell(row=linha, column=i, value=v)
    linha += 2
    ws.cell(
        row=linha, column=1, value="Viés do juiz (mesma amostra): nota de cada juiz por modelo julgado"
    ).font = negrito
    linha += 1
    cabecalho(
        ws,
        linha,
        [
            "Modelo julgado",
            "n",
            "Utilidade Opus",
            "Utilidade GLM",
            "Exatidão Opus",
            "Exatidão GLM",
            "Naturalidade Opus",
            "Naturalidade GLM",
            "% seguiu Opus",
            "% seguiu GLM",
        ],
    )
    for m, x in conc["por_modelo"].items():
        linha += 1
        for i, v in enumerate(
            [
                m,
                x["n"],
                x["utilidade_opus"],
                x["utilidade_glm"],
                x["exatidao_opus"],
                x["exatidao_glm"],
                x["naturalidade_opus"],
                x["naturalidade_glm"],
                x["seguiu_opus"],
                x["seguiu_glm"],
            ],
            start=1,
        ):
            ws.cell(row=linha, column=i, value=v)
    linha += 2
    ws.cell(row=linha, column=1, value="Observações de método").font = negrito
    notas = [
        "Viés: Claude Opus 5 julgando Claude Sonnet/Haiku pode favorecê-los; por isso o GLM-5 julgou de novo "
        "uma amostra (tabela acima). Diferença sistemática entre as colunas Opus/GLM do mesmo modelo = viés.",
        "Temperatura 0,2 vale para Mistral e GLM-5. Os Claude 5.x recusam o parâmetro `temperature` "
        "(ValidationException) e rodaram sem ele, com esforço baixo (output_config.effort=low).",
        "Mesmo prompt de sistema, mesmo max_tokens (2000) para os 4 modelos. Latência medida dentro da AWS "
        "(tarefa ECS em sa-east-1); Claude pelo perfil GLOBAL (pode processar fora do Brasil).",
        "Custo: preço on-demand do Bedrock sa-east-1 (API de preços da AWS, 17/10/2026), entrada+saída de todas "
        "as chamadas de geração (inclui reescritas); não inclui o juiz.",
        "A 'resposta atual da Helena' das 9 perguntas do dono é RECONSTRUÍDA com as mesmas funções do repo "
        "(busca da base, textos fixos, perguntas de alerta), não pelo grafo inteiro.",
        "A verificação automática não conhece o comportamento esperado: usa as cercas do repo e os detectores "
        "de alarme/risco que a Helena já tem sobre a mensagem.",
    ]
    for n in notas:
        linha += 1
        ws.cell(row=linha, column=1, value=n).alignment = quebra
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=6)
        ws.row_dimensions[linha].height = 32
    larguras(ws, {1: 52, 2: 18, 3: 18, 4: 18, 5: 18, 6: 18, 7: 18, 8: 18, 9: 14, 10: 14})

    # --- Por categoria ---
    ws = wb.create_sheet("Por categoria")
    cols = ["Categoria", "n"]
    for m in modelos:
        cols += [f"{m} utilidade", f"{m} exatidão", f"{m} % seguiu", f"{m} % conduta indevida"]
    cabecalho(ws, 1, cols)
    for i, (cat, por_m) in enumerate(sorted(resumo["por_categoria"].items()), start=2):
        ws.cell(row=i, column=1, value=cat)
        ws.cell(row=i, column=2, value=max(x["n"] for x in por_m.values()))
        col = 3
        for m in modelos:
            x = por_m.get(m, {})
            for k in ("media_utilidade", "media_exatidao", "pct_seguiu_esperado", "pct_conduta_indevida"):
                ws.cell(row=i, column=col, value=x.get(k))
                col += 1
    larguras(ws, {1: 26, 2: 6, **{c: 12 for c in range(3, 3 + 4 * len(modelos))}})
    ws.freeze_panes = "C2"

    # --- Respostas ---
    ws = wb.create_sheet("Respostas")
    cabecalho(
        ws,
        1,
        [
            "id",
            "categoria",
            "esperado",
            "pergunta (e histórico)",
            "modelo",
            "resposta enviada",
            "desfecho da verificação",
            "o que a verificação apontou (1a tentativa)",
            "apoio injetado",
            "utilidade",
            "exatidão",
            "naturalidade",
            "acertou assunto",
            "seguiu esperado",
            "conduta indevida",
            "qual conduta",
            "usou o que disse",
            "tamanho ok",
            "fonte",
            "erros factuais (juiz)",
            "justificativa do juiz",
        ],
    )
    linha = 1
    for p in perguntas:
        texto_p = p.texto if not p.historico else f"[antes: {' | '.join(p.historico)}] {p.texto}"
        for m in modelos:
            r = resp.get(f"{p.id}|{m}")
            if r is None:
                continue
            v = j1.get(f"{p.id}|{m}", {})
            linha += 1
            viol = "; ".join(
                f"{x['regra']}" + (f" ({x['trecho']})" if x.get("trecho") else "")
                for x in r.get("violacoes1", [])
            )
            apoio = r.get("apoio") or {}
            apoio_txt = (
                (
                    f"{apoio.get('assunto_id')} ({apoio.get('como')})"
                    + (f" — fonte: {apoio['fonte_nome']}" if apoio.get("fonte_nome") else "")
                )
                if apoio
                else ""
            )
            valores = [
                p.id,
                p.categoria,
                p.esperado,
                texto_p,
                m,
                r.get("final") or f"ERRO: {r.get('erro')}",
                r.get("desfecho"),
                viol,
                apoio_txt,
                v.get("utilidade"),
                v.get("exatidao"),
                v.get("naturalidade"),
                v.get("acertou_assunto"),
                v.get("seguiu_esperado"),
                v.get("conduta_indevida"),
                v.get("conduta_indevida_qual"),
                v.get("usou_o_que_disse"),
                v.get("tamanho_adequado"),
                v.get("fonte"),
                "; ".join(v.get("erros_factuais") or []),
                v.get("justificativa") or v.get("erro"),
            ]
            for i, val in enumerate(valores, start=1):
                ws.cell(row=linha, column=i, value=val).alignment = quebra
    larguras(
        ws,
        {
            1: 6,
            2: 16,
            3: 18,
            4: 34,
            5: 16,
            6: 70,
            7: 11,
            8: 30,
            9: 24,
            10: 8,
            11: 8,
            12: 8,
            13: 8,
            14: 8,
            15: 8,
            16: 24,
            17: 10,
            18: 8,
            19: 11,
            20: 34,
            21: 46,
        },
    )
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:U{linha}"

    # --- Amostra para o dono ---
    ws = wb.create_sheet("Amostra para o dono")
    ws["A1"] = (
        "30 perguntas cobrindo as categorias. As 4 respostas estão EMBARALHADAS em cada linha (A-D) para "
        "não influenciar; a chave está nas colunas ocultas à direita. Marque a melhor na coluna H."
    )
    ws["A1"].alignment = quebra
    ws.merge_cells("A1:H1")
    ws.row_dimensions[1].height = 30
    cabecalho(
        ws,
        2,
        [
            "id",
            "categoria",
            "pergunta",
            "Resposta A",
            "Resposta B",
            "Resposta C",
            "Resposta D",
            "Melhor (A/B/C/D) e comentário",
            "chave A",
            "chave B",
            "chave C",
            "chave D",
        ],
    )
    sorteio = random.Random(20261017)
    for i, p in enumerate(amostra(perguntas), start=3):
        ordem = list(modelos)
        sorteio.shuffle(ordem)
        texto_p = p.texto if not p.historico else f"[antes: {' | '.join(p.historico)}] {p.texto}"
        valores = [
            p.id,
            p.categoria,
            texto_p,
            *[(resp.get(f"{p.id}|{m}") or {}).get("final") for m in ordem],
            None,
            *ordem,
        ]
        for c, val in enumerate(valores, start=1):
            ws.cell(row=i, column=c, value=val).alignment = quebra
    larguras(ws, {1: 6, 2: 16, 3: 30, 4: 48, 5: 48, 6: 48, 7: 48, 8: 24})
    for col in "IJKL":
        ws.column_dimensions[col].hidden = True
    ws.freeze_panes = "D3"

    # --- Perguntas do dono ---
    ws = wb.create_sheet("Perguntas do dono")
    cols = [
        "id",
        "pergunta",
        "resposta ATUAL da Helena (reconstruída)",
        "como a atual chegou lá",
        "atual: utilidade / exatidão / seguiu (juiz)",
    ]
    for m in modelos:
        cols += [f"{m}: resposta", f"{m}: utilidade / exatidão / seguiu"]
    cabecalho(ws, 1, cols)

    def notas(v: dict[str, Any]) -> str:
        if not v or "erro" in v:
            return ""
        return f"{v['utilidade']} / {v['exatidao']} / {'sim' if v['seguiu_esperado'] else 'não'} — {v['justificativa']}"

    for i, p in enumerate([p for p in perguntas if p.do_dono], start=2):
        atual = d["atual"].get(p.id, {})
        texto_p = p.texto if not p.historico else f"[antes: {' | '.join(p.historico)}] {p.texto}"
        valores = [p.id, texto_p, atual.get("texto"), atual.get("como"), notas(j1.get(f"{p.id}|atual", {}))]
        for m in modelos:
            valores += [(resp.get(f"{p.id}|{m}") or {}).get("final"), notas(j1.get(f"{p.id}|{m}", {}))]
        for c, val in enumerate(valores, start=1):
            ws.cell(row=i, column=c, value=val).alignment = quebra
    larguras(ws, {1: 6, 2: 28, 3: 50, 4: 22, 5: 30, **{c: 44 for c in range(6, 6 + 2 * len(modelos))}})
    ws.freeze_panes = "C2"

    wb.save(xlsx)
    atual_vs = {}
    donos = [p.id for p in perguntas if p.do_dono]
    for nome in ["atual", *modelos]:
        vs = [j1.get(f"{pid}|{nome}") for pid in donos]
        vs = [v for v in vs if v and "erro" not in v]
        atual_vs[nome] = {
            "n": len(vs),
            "utilidade": round(sum(v["utilidade"] for v in vs) / len(vs), 2) if vs else None,
            "exatidao": round(sum(v["exatidao"] for v in vs) / len(vs), 2) if vs else None,
            "seguiu": sum(1 for v in vs if v["seguiu_esperado"]),
        }
    resumo["perguntas_do_dono"] = atual_vs
    return resumo


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="planilha do experimento redacao de saude")
    p.add_argument("--resultados", required=True)
    p.add_argument("--perguntas", required=True)
    p.add_argument("--xlsx", required=True)
    p.add_argument("--resumo-json", default="", help="grava tambem o resumo calculado (para o relatorio)")
    a = p.parse_args(argv)
    resumo = montar(Path(a.resultados), Path(a.perguntas), Path(a.xlsx))
    if a.resumo_json:
        Path(a.resumo_json).write_text(json.dumps(resumo, ensure_ascii=False, indent=2), encoding="utf-8")
    print(a.xlsx)
    return 0


if __name__ == "__main__":
    sys.exit(main())
